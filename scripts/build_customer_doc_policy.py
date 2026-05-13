# -*- coding: utf-8 -*-
"""Build a unified Excel file for customer doc-policy (per-order / aggregated
delivery notes & invoices) plus a separate sheet listing customers with
missing city/zone so they can be fixed in SAP.

Source: the two xlsx files the user uploaded in his Temp folder.
Output: docs/customer-doc-policy/customer-doc-policy.xlsx
"""
import os
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

OIG = r'C:\Users\izik\AppData\Local\Temp\של רשימה של כרטיסיםOIG (3).xlsx'
UNI = r'C:\Users\izik\AppData\Local\Temp\רשימה של כרטיסים UNICO (3).xlsx'
OUT_DIR = r'C:\Users\izik\OneDrive - OIG\שולחן העבודה\cowork\sap-logistics-hub\docs\customer-doc-policy'
os.makedirs(OUT_DIR, exist_ok=True)
OUT = os.path.join(OUT_DIR, 'customer-doc-policy.xlsx')

oig = pd.read_excel(OIG, sheet_name='Sheet1')
uni = pd.read_excel(UNI, sheet_name='גיליון1')

print('OIG rows:', len(oig), '| UNICO rows:', len(uni))


def norm_oig(r):
    return {
        'חברה': 'OIG',
        'קוד כרטיס': r['קוד כרטיס'],
        'שם כרטיס': r['שם כרטיס'],
        'עיר': r.get('כתובת לחיוב - עיר', ''),
        'רחוב': r.get('כתובת לחיוב - רחוב', ''),
        'אזור הפצה': r.get('אזור', ''),
        'יום הפצה 1': r.get('ימי הפצה', ''),
        'יום הפצה 2': r.get('ימי הפצה.1', ''),
    }


def norm_uni(r):
    city = r.get('כתובת למשלוח - עיר') or r.get('כתובת לחיוב - עיר') or ''
    street = r.get('כתובת למשלוח - רחוב') or r.get('רחוב') or ''
    return {
        'חברה': 'UNICO',
        'קוד כרטיס': r['קוד כרטיס'],
        'שם כרטיס': r['שם כרטיס'],
        'עיר': city,
        'רחוב': street,
        'אזור הפצה': r.get('אזור', ''),
        'יום הפצה 1': r.get('יום אספקה', ''),
        'יום הפצה 2': r.get('יום אספקה ', ''),
    }


rows = [norm_oig(r) for _, r in oig.iterrows()] + [norm_uni(r) for _, r in uni.iterrows()]
df = pd.DataFrame(rows)
df['תעודת משלוח לכל הזמנה'] = ''
df['חשבונית לכל הזמנה'] = ''
df['תעודת משלוח מרכזת לקו'] = ''
df['חשבונית מרכזת לקו'] = ''
df['הערות'] = ''


def issue_reason(row):
    city = str(row['עיר']).strip().lower()
    zone = str(row['אזור הפצה']).strip().lower()
    issues = []
    if not city or city in ('nan', 'none'):
        issues.append('חסר עיר')
    if not zone or zone in ('nan', 'none'):
        issues.append('חסר אזור')
    return ' + '.join(issues) if issues else ''


df['בעיית שיוך'] = df.apply(issue_reason, axis=1)
issues_df = df[df['בעיית שיוך'] != ''].copy()
clean_df = df[df['בעיית שיוך'] == ''].copy().drop(columns=['בעיית שיוך'])

print('Clean:', len(clean_df), '| Issues:', len(issues_df))

wb = Workbook()


def write_sheet(ws, frame, with_dropdowns=True):
    cols = list(frame.columns)
    header_fill = PatternFill('solid', start_color='2563eb')
    header_font = Font(bold=True, color='FFFFFF', name='Arial', size=11)
    thin = Side(border_style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    body_font = Font(name='Arial', size=10)
    for r, row in enumerate(frame.itertuples(index=False), 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val if pd.notna(val) else '')
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(horizontal='right', vertical='center')

    ws.freeze_panes = 'A2'
    ws.sheet_view.rightToLeft = True

    widths = {
        'חברה': 8, 'קוד כרטיס': 12, 'שם כרטיס': 45,
        'עיר': 16, 'רחוב': 28, 'אזור הפצה': 14,
        'יום הפצה 1': 12, 'יום הפצה 2': 12,
        'תעודת משלוח לכל הזמנה': 16, 'חשבונית לכל הזמנה': 16,
        'תעודת משלוח מרכזת לקו': 18, 'חשבונית מרכזת לקו': 18,
        'הערות': 22, 'בעיית שיוך': 18,
    }
    for c, name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 14)

    if with_dropdowns:
        dv = DataValidation(type='list', formula1='"כן,לא,—"', allow_blank=True)
        dv.error = 'בחר: כן / לא / —'
        dv.errorTitle = 'ערך לא חוקי'
        dv.prompt = 'בחר: כן / לא / —'
        dv.promptTitle = 'מדיניות מסמך'
        ws.add_data_validation(dv)
        for col_name in ['תעודת משלוח לכל הזמנה', 'חשבונית לכל הזמנה', 'תעודת משלוח מרכזת לקו', 'חשבונית מרכזת לקו']:
            if col_name in cols:
                col_idx = cols.index(col_name) + 1
                letter = get_column_letter(col_idx)
                dv.add('{0}2:{0}{1}'.format(letter, len(frame) + 1))


ws1 = wb.active
ws1.title = 'לקוחות תקינים'
write_sheet(ws1, clean_df, with_dropdowns=True)

ws2 = wb.create_sheet('בעיית שיוך - לעדכן SAP')
issue_cols_order = ['חברה', 'קוד כרטיס', 'שם כרטיס', 'עיר', 'רחוב', 'אזור הפצה', 'בעיית שיוך', 'יום הפצה 1', 'יום הפצה 2']
write_sheet(ws2, issues_df[issue_cols_order], with_dropdowns=False)

ws3 = wb.create_sheet('הסבר')
ws3.sheet_view.rightToLeft = True
ws3['A1'] = 'הסבר — טבלת מדיניות מסמכים ללקוח'
ws3['A1'].font = Font(bold=True, size=14, name='Arial')
inst = [
    '',
    'הקובץ כולל 2 גיליונות עבודה:',
    '  1. לקוחות תקינים — לקוחות עם עיר + אזור מוגדרים. למלא 4 העמודות:',
    '       • תעודת משלוח לכל הזמנה — כן/לא',
    '       • חשבונית לכל הזמנה — כן/לא',
    '       • תעודת משלוח מרכזת לקו — כן/לא',
    '       • חשבונית מרכזת לקו — כן/לא',
    '       • הערות — חופשי',
    '',
    '  2. בעיית שיוך - לעדכן SAP — לקוחות ללא עיר ו/או ללא אזור.',
    '       יש לעדכן את הכרטיס ב-SAP בעמודה הרלוונטית (כתובת חיוב / משלוח / אזור)',
    '       ואז להריץ את הסקריפט שוב כדי לרענן.',
    '',
    'ערכים מותרים בעמודות הבחירה: כן / לא / — (לא רלוונטי).',
    '',
    'הקובץ נוצר אוטומטית מ-2 הקבצים שהעלית:',
    '   • של רשימה של כרטיסיםOIG (3).xlsx',
    '   • רשימה של כרטיסים UNICO (3).xlsx',
    '',
    'סה"כ {0} כרטיסים. תקינים: {1}. עם בעיית שיוך: {2}.'.format(len(df), len(clean_df), len(issues_df)),
    '',
    'תאריך יצירה: ' + pd.Timestamp.now().strftime('%Y-%m-%d %H:%M'),
]
for i, line in enumerate(inst, start=2):
    cell = ws3.cell(row=i, column=1, value=line)
    cell.font = Font(name='Arial', size=11)
    cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
ws3.column_dimensions['A'].width = 100

wb._sheets = [ws1, ws2, ws3]
wb.save(OUT)

print('Wrote:', OUT)
print('  לקוחות תקינים:', len(clean_df))
print('  בעיית שיוך:', len(issues_df))
print('  file size:', os.path.getsize(OUT), 'bytes')
